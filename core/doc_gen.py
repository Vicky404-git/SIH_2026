# doc_gen_tool.py
import os
import subprocess
import tempfile
from .agent import Tool, ToolResult
from .config import get_sandbox_limits
from .sandbox import _apply_os_limits

def generate_docx(markdown_content: str) -> ToolResult:
    limits = get_sandbox_limits()

    with tempfile.NamedTemporaryFile(suffix=".md", delete=False, mode="w", encoding="utf-8") as f:
        f.write(markdown_content)
        md_path = f.name
    out_path = md_path.replace(".md", ".docx")

    def _pandoc_limits():
        _apply_os_limits(max_mem_mb=limits["max_mem_mb"], max_cpu_sec=limits["timeout_sec"])

    # Try pandoc first if installed
    try:
        result = subprocess.run(
            ["pandoc", md_path, "-o", out_path],
            preexec_fn=_pandoc_limits if os.name != "nt" else None,
            capture_output=True,
            text=True,
            timeout=limits["timeout_sec"],
            shell=(os.name == "nt"),
        )
        if result.returncode == 0 and os.path.exists(out_path):
            return ToolResult(ok=True, output=out_path)
    except Exception:
        pass

    # Reliable python-docx fallback (works 100% offline without external CLI tools)
    try:
        from docx import Document
        doc = Document()
        for block in markdown_content.split("\n\n"):
            b = block.strip()
            if not b:
                continue
            if b.startswith("# "):
                doc.add_heading(b[2:], level=1)
            elif b.startswith("## "):
                doc.add_heading(b[3:], level=2)
            elif b.startswith("### "):
                doc.add_heading(b[4:], level=3)
            elif b.startswith("- ") or b.startswith("* "):
                for line in b.split("\n"):
                    doc.add_paragraph(line.lstrip("- *"), style='List Bullet')
            else:
                doc.add_paragraph(b)
        doc.save(out_path)
        return ToolResult(ok=True, output=out_path)
    except Exception as e:
        return ToolResult(ok=False, output=f"Failed to generate docx: {e}")

docgen_tool = Tool(
    "generate_docx", 
    generate_docx, 
    "ONLY use this tool if the user explicitly asks to 'generate a document', 'create a report', or 'make a .docx file'. Do NOT use this for code."
)
